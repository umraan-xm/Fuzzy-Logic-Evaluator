import sys
import PyQt5
import openpyxl
from PyQt5.QtWidgets import *
from PyQt5.QtGui import *
from PyQt5.QtCore import *
from PyQt5.QtSql import QSqlDatabase, QSqlQuery
import pandas
import pandas as pd
import Student
import Result
import Delegates
import Fuzzy
import Guidelines


class MainWindow(QMainWindow):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Fuzzy Evaluator")
        self.resize(1280, 720)
        self.setWindowIcon(QIcon('project_icon.jpg'))

        #self.showMaximized()

        self.centralwidget = QWidget(self)
        self.setCentralWidget(self.centralwidget)

        # ---------------------------  MENU BAR ---------------------------

        self.actionNew = QAction("New", self)
        self.actionNew.setShortcut('Ctrl+N')
        self.actionNew.triggered.connect(self.new_file)

        self.actionSave = QAction("Save", self)
        self.actionSave.setShortcut('Ctrl+S')
        self.actionSave.triggered.connect(self.save_file)

        self.actionSaveAs = QAction("Save As...", self)
        self.actionSaveAs.setShortcut('Ctrl+Shift+S')
        self.actionSaveAs.triggered.connect(self.save_file_as)

        self.actionDelete = QAction("Delete", self)
        self.actionDelete.setShortcut('Delete')
        self.actionDelete.triggered.connect(self.delete_file)

        self.actionReset = QAction("Reset", self)
        self.actionReset.triggered.connect(self.reset)

        self.actionGuidelines = QAction("Guidelines", self)
        self.actionGuidelines.triggered.connect(self.guidelines)

        self.actionGroupTheme = QActionGroup(self)
        self.actionDark = QAction("Dark", self)
        self.actionDark.setCheckable(True)
        self.actionGroupTheme.addAction(self.actionDark)
        self.actionLight = QAction("Light", self)
        self.actionLight.setCheckable(True)
        self.actionGroupTheme.addAction(self.actionLight)
        self.actionDark.setChecked(True)
        self.actionGroupTheme.triggered.connect(self.toggleTheme)

        self.menubar = self.menuBar()
        self.menubar.setObjectName("menubar")

        self.menuFile = QMenu()
        self.menuFile.setObjectName("menuFile")
        self.menuFile.setTitle("File")
        self.menubar.addMenu(self.menuFile)
        self.menuFile.addAction(self.actionNew)
        self.menuFile.addAction(self.actionSave)
        self.menuFile.addAction(self.actionSaveAs)
        self.menuFile.addAction(self.actionDelete)
        self.menuFile.addSeparator()
        self.menuFile.addAction(self.actionReset)

        self.menuTheme = QMenu()
        self.menuTheme.setObjectName("menuTheme")
        self.menuTheme.setTitle("Theme")
        self.menubar.addMenu(self.menuTheme)
        self.menuTheme.addAction(self.actionDark)
        self.menuTheme.addAction(self.actionLight)

        self.menuHelp = QMenu()
        self.menuHelp.setObjectName("menuHelp")
        self.menuHelp.setTitle("Help")
        self.menubar.addMenu(self.menuHelp)
        self.menuHelp.addAction(self.actionGuidelines)

        # ---------------------------  MENU BAR END / TOOL BAR ---------------------------

        self.toolbar = QToolBar("File", self)
        self.toolbar.setMovable(False)

        self.fileList = QListWidget()
        self.fileList.addItem("Untitled")

        self.toolbar.addWidget(self.fileList)

        self.addToolBar(Qt.LeftToolBarArea, self.toolbar)

        # ---------------------------  TOOL BAR END / DOCK WIDGET ---------------------------

        self.dockWidget = QDockWidget("Tools", self)
        self.dockWidget.setAllowedAreas(Qt.RightDockWidgetArea)
        self.dockWidget.setFeatures(QDockWidget.DockWidgetMovable | QDockWidget.DockWidgetFloatable)

        self.dockLayout = QVBoxLayout()

        self.dockMultiWidget = QWidget()
        # Compute button
        self.btnCompute = QPushButton("Compute")
        # Clear button
        self.btnClear = QPushButton("Clear")
        # Create button
        self.btnCreateTable = QPushButton("Create Database Table")
        # Save button
        self.btnSave = QPushButton("Save")
        # Delete button
        self.btnDelete = QPushButton("Delete")
        # Load excel sheet
        self.btnLoad = QPushButton("Load Student Marks")
        # Export result
        self.btnExport = QPushButton("Export Results")
        # Export all
        self.btnExportAll = QPushButton("Export All")

        # Start roll spinbox
        self.lblStartRoll = QLabel("&Starting Roll No: ")
        self.spStartRoll = QSpinBox()
        self.spStartRoll.setMinimum(1)
        self.lblStartRoll.setBuddy(self.spStartRoll)
        self.startRollWidget = QWidget()
        self.layoutStartRoll = QHBoxLayout()
        self.layoutStartRoll.addWidget(self.lblStartRoll)
        self.layoutStartRoll.addWidget(self.spStartRoll)
        self.startRollWidget.setLayout(self.layoutStartRoll)

        # End roll spinbox
        self.lblEndRoll = QLabel("&Last Roll No: ")
        self.spEndRoll = QSpinBox()
        self.spEndRoll.setMinimum(1)
        self.spEndRoll.setMaximum(1000000)
        self.spEndRoll.setValue(100)
        self.spStartRoll.setMaximum(self.spEndRoll.value())
        self.lblEndRoll.setBuddy(self.spEndRoll)
        self.endRollWidget = QWidget()
        self.layoutEndRoll = QHBoxLayout()
        self.layoutEndRoll.addWidget(self.lblEndRoll)
        self.layoutEndRoll.addWidget(self.spEndRoll)
        self.endRollWidget.setLayout(self.layoutEndRoll)

        self.dockLayout.addWidget(self.startRollWidget)
        self.dockLayout.addWidget(self.endRollWidget)
        self.dockLayout.addWidget(self.btnCompute)
        self.dockLayout.addWidget(self.btnClear)
        self.dockLayout.addStretch()
        self.dockLayout.addWidget(QLabel("Excel Controls: "))
        self.dockLayout.addWidget(self.btnLoad)
        self.dockLayout.addWidget(self.btnExport)
        self.dockLayout.addWidget(self.btnExportAll)
        self.dockLayout.addStretch()
        self.dockLayout.addWidget(QLabel("Database Controls: "))
        self.dockLayout.addWidget(self.btnCreateTable)
        self.dockLayout.addWidget(self.btnSave)
        self.dockLayout.addWidget(self.btnDelete)
        self.dockLayout.addStretch()
        self.dockMultiWidget.setLayout(self.dockLayout)

        self.dockWidget.setWidget(self.dockMultiWidget)
        self.addDockWidget(Qt.RightDockWidgetArea, self.dockWidget)

        # ---------------------------  DOCK WIDGET END / STATUS BAR ---------------------------

        self.status = self.statusBar()

        # ---------------------------  STATUS BAR END / TABLES ---------------------------

        self.tableStudent = QTableView(self)
        self.tableStudentModel = Student.TableStudentModel()
        self.tableStudent.setModel(self.tableStudentModel)
        self.tableStudent.setItemDelegateForColumn(0, Delegates.StringItemDelegate(self))
        self.tableStudent.setItemDelegateForColumn(1, Delegates.IntegerItemDelegate(self))
        self.tableStudent.setItemDelegateForColumn(2, Delegates.IntegerItemDelegate(self))
        self.tableStudent.setItemDelegateForColumn(3, Delegates.IntegerItemDelegate(self))
        self.tableStudent.setEditTriggers(QAbstractItemView.CurrentChanged)
        # self.tableStudent.setSortingEnabled(True)

        self.setTableStudentUI()

        self.tableResult = QTableView(self)
        self.tableResultModel = Result.TableResultModel()
        self.tableResult.setModel(self.tableResultModel)
        self.setTableResultUI()

        # -------------------- TABLE END --------------------------------------------------

        self.centralLayout = QHBoxLayout(self.centralwidget)
        self.centralLayout.addWidget(self.tableStudent)
        self.centralLayout.addWidget(self.tableResult)
        self.centralLayout.setStretch(0, 100)

        self.btnCompute.clicked.connect(self.compute)
        self.btnClear.clicked.connect(self.clear)
        self.spEndRoll.valueChanged.connect(self.updateRollNo)
        self.spStartRoll.valueChanged.connect(self.updateRollNo)
        self.btnCreateTable.clicked.connect(self.save_file_as)
        self.btnSave.clicked.connect(self.save_file)
        self.btnDelete.clicked.connect(self.delete_file)
        self.btnLoad.clicked.connect(self.load_excel)
        self.btnExport.clicked.connect(self.export_result)
        self.btnExportAll.clicked.connect(self.export_all)
        self.tableStudentModel.dataChanged.connect(self.unsaved)
        self.fileList.itemClicked.connect(self.load)

        self.setup_database()
        self.load_all_tables()

        self.toggleTheme()

    def load(self):
        filename = self.fileList.currentItem().text()

        if "*" in filename:
            filename = filename.replace("*", "")

        if filename == 'Untitled':
            self.tableStudentModel.clear()
            self.tableStudent.setModel(None)
            self.tableStudent.setModel(self.tableStudentModel)
            self.spStartRoll.setValue(1)
            self.spEndRoll.setValue(100)
            self.setTableStudentUI()

            self.tableResultModel.clear()
            self.tableResult.setModel(None)
            self.tableResult.setModel(self.tableResultModel)
            self.setTableResultUI()
            return

        queryMinRoll = QSqlQuery()
        queryMinRoll.exec(f"SELECT min(id) from {filename}")
        queryMaxRoll = QSqlQuery()
        queryMaxRoll.exec(f"SELECT max(id) from {filename}")
        if queryMinRoll.next() and queryMaxRoll.next():
            start_roll = queryMinRoll.value(0)
            end_roll = queryMaxRoll.value(0)

        self.spEndRoll.setValue(end_roll+1)
        self.spStartRoll.setValue(start_roll)

        querySelectStudent = QSqlQuery()
        querySelectStudent.exec(f"SELECT name, exam1, exam2, practicals FROM {filename}")
        records = []
        while querySelectStudent.next():
            records.append([querySelectStudent.value(0), str(querySelectStudent.value(1)), str(querySelectStudent.value(2)), str(querySelectStudent.value(3))])
        self.tableStudentModel.setAllRecords(records)

        querySelectResults = QSqlQuery()
        querySelectResults.exec(f"SELECT result FROM {filename}")
        results = []
        while querySelectResults.next():
            results.append(str(querySelectResults.value(0)))
        self.tableResultModel.setAllRecords(results)

    def load_all_tables(self):
        table_list = self.con.tables()
        if 'sqlite_sequence' in table_list:
            table_list.remove('sqlite_sequence')
            if table_list:
                for table in table_list:
                    self.fileList.addItem(table)

    def unsaved(self):
        current_item = self.fileList.currentItem()
        if current_item is not None:
            current_text = current_item.text()
            if current_text[-1] != '*':
                self.fileList.currentItem().setText(current_text + "*")

    def new_file(self):
        self.fileList.addItem("Untitled")
        new_item = self.fileList.item(self.fileList.count() - 1)
        self.fileList.setCurrentItem(new_item)
        self.load()

    def save_file(self):
        try:
            filename = self.fileList.currentItem().text()

            if "*" in filename:
                filename = filename.replace("*", "")

            if filename not in self.con.tables():
                answer = QMessageBox.question(self, '', "This table does not exist in the database. Would you like to create one?",
                                     QMessageBox.Yes | QMessageBox.No)
                if answer == QMessageBox.Yes:
                    self.save_file_as()
                else:
                    return

            self.fileList.currentItem().setText(filename)
            data = self.tableStudentModel.getAllRecords()
            result = self.tableResultModel.getAllRecords()
            start_roll = int(self.spStartRoll.value())

            while len(result) < len(data):
                result.append(0)
            for record, res in zip(data, result):
                record.append(res)

            queryMaxRoll = QSqlQuery()
            queryMaxRoll.exec(f"SELECT max(id) from {filename}")
            queryMaxRoll.next()
            max_roll = queryMaxRoll.value(0)

            for record in data:
                if start_roll <= max_roll:
                    start_roll += 1
                    continue
                name = record[0]
                exam1 = round(float(record[1]), 2)
                exam2 = round(float(record[2]), 2)
                practicals = round(float(record[3]), 2)
                final_result = round(float(record[4]), 2)

                queryInsertData = QSqlQuery()
                queryInsertData.exec(
                    f"""INSERT INTO {filename} VALUES
                    ('{start_roll}', '{name}', '{exam1}', '{exam2}', '{practicals}', '{final_result}')
                                """
                )

                # print(start_roll, " ", name, " ", exam1, " ", exam2, " ", practicals, " ", final_result)
                start_roll += 1

            self.status.showMessage("Records successfully saved to database.", 2000)
        except Exception as e:
            QMessageBox.warning(self, "Error", str(e), QMessageBox.Ok)

    def save_file_as(self):
        filename, ok = QInputDialog().getText(self, "Save As", "Enter the name of the Table: ")
        rx = QRegExp("[A-Za-z][A-Za-z0-9_]*")
        if not rx.exactMatch(filename):
            QMessageBox.warning(self, "Invalid",
                                "Invalid name for table. Only letters, digits and underscores allowed. Do not start the name with a digit. ",
                                QMessageBox.Ok)
            return

        if ok:
            self.fileList.currentItem().setText(filename)

            queryCreateTable = QSqlQuery()
            queryCreateTable.exec(
                f"""
                CREATE TABLE IF NOT EXISTS {filename} (
                    id INTEGER PRIMARY KEY AUTOINCREMENT UNIQUE NOT NULL,
                    name VARCHAR(100) NOT NULL,
                    exam1 REAL NOT NULL,
                    exam2 REAL NOT NULL,
                    practicals REAL NOT NULL,
                    result REAL
                )
                """
            )
            # print(self.con.tables())

            data = self.tableStudentModel.getAllRecords()
            result = self.tableResultModel.getAllRecords()
            start_roll = int(self.spStartRoll.value())

            while len(result) < len(data):
                result.append(0)
            for record, res in zip(data, result):
                record.append(res)

            for record in data:
                name = record[0]
                exam1 = round(float(record[1]), 2)
                exam2 = round(float(record[2]), 2)
                practicals = round(float(record[3]), 2)
                final_result = round(float(record[4]), 2)

                queryInsertData = QSqlQuery()
                queryInsertData.exec(
                    f"""INSERT INTO {filename}
                        VALUES ('{start_roll}', '{name}', '{exam1}', '{exam2}', '{practicals}', '{final_result}')
                    """
                )

                # print(start_roll, " ", name, " ", exam1, " ", exam2, " ", practicals, " ", final_result)
                start_roll += 1

            self.status.showMessage("Table " + filename + " created successfully.", 2000)

    def delete_file(self):
        filename = self.fileList.currentItem().text()

        if "*" in filename:
            filename = filename.replace("*", "")

        if filename not in self.con.tables():
            QMessageBox.warning(self, "Error", "This table does not exist in the database!", QMessageBox.Ok)
            return

        answer = QMessageBox.question(self, '', "Are you sure you want to delete this table?", QMessageBox.Yes | QMessageBox.No)
        if answer == QMessageBox.Yes:
            queryDeleteTable = QSqlQuery()
            queryDeleteTable.exec(
                f"DROP TABLE IF EXISTS {filename}"
            )
            temp = self.fileList.item(0)
            self.fileList.setCurrentItem(temp)
            items_list = self.fileList.findItems(filename, Qt.MatchExactly)
            row = self.fileList.row(items_list[0])
            self.fileList.takeItem(row)
            self.load()
            self.status.showMessage("Table " + filename + " deleted successfully.", 2000)
        else:
            return

    def reset(self):
        answer = QMessageBox.question(self, '', "Are you sure you want to delete all tables?",
                                      QMessageBox.Yes | QMessageBox.No)

        if answer == QMessageBox.Yes:
            table_list = self.con.tables()
            if 'sqlite_sequence' in table_list:
                table_list.remove('sqlite_sequence')

            for table in table_list:
                queryDeleteTable = QSqlQuery()
                queryDeleteTable.exec(
                    f"DROP TABLE IF EXISTS {table}"
                )

            temp = self.fileList.item(0)
            self.fileList.setCurrentItem(temp)
            items = [self.fileList.item(x) for x in range(self.fileList.count())]
            for item in items:
                if (item.text() == "Untitled") or (item.text() == "Untitled*"):
                    continue
                row = self.fileList.row(item)
                self.fileList.takeItem(row)
            self.load()

            self.status.showMessage("All tables have been deleted successfully.", 2000)
        else:
            return

    def load_excel(self):
        try:
            file = QFileDialog.getOpenFileUrl(self, "Open Excel sheet", filter="Excel File (*.xlsx *.xls)")
            if not file[1]:
                return
            file = file[0].path()[1:]
            data = pd.read_excel(file).values.tolist()

            if len(data[0]) == 5:
                start = data[0][0]
                end = data[-1][0]
                for row in data:
                    row.pop(0)
                self.spEndRoll.setValue(end)
                self.spStartRoll.setValue(start)
            for record in data:
                if len(record) != 4:
                    QMessageBox.warning(self, "Error", "Invalid excel sheet. Please make sure there are only 4 columns."
                                                       "If there are 5 columns, please make sure the first column "
                                                       "contains the roll numbers in ascending order", QMessageBox.Ok)
                    return

            for i in range(len(data)):
                record = [str(x) for x in data[i]]
                data[i] = record
            self.tableStudentModel.setAllRecords(data)

            self.status.showMessage("Excel data loaded successfully", 2000)
        except Exception as e:
            QMessageBox.warning(self, "Error", "Invalid excel sheet. Please make sure there are only 4 columns."
                                               "If there are 5 columns, please make sure the first column "
                                               "contains the roll numbers in ascending order", QMessageBox.Ok)

    def export_result(self):
        file = QFileDialog.getSaveFileName(self, "Save File", filter="Excel File (*.xlsx)")
        if not file[1]:
            return
        file = file[0]

        result = self.tableResultModel.getAllRecords()
        start = self.spStartRoll.value()
        end = self.spEndRoll.value()
        roll = [i for i in range(start, end)]
        # records = [list(x) for x in list(zip(roll, result))]
        records = []
        for i in range(len(result)):
            records.append([roll[i], float(result[i])])
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Roll No.", "Results"])
            for record in records:
                ws.append(record)
            wb.save(file)
            self.status.showMessage("Results successfully saved to " + file, 2000)
        except Exception as e:
            QMessageBox.warning(self, "Error", str(e), QMessageBox.Ok)

    def export_all(self):
        file = QFileDialog.getSaveFileName(self, "Save File", filter="Excel File (*.xlsx)")
        if not file[1]:
            return
        file = file[0]

        data = self.tableStudentModel.getAllRecords()
        result = self.tableResultModel.getAllRecords()
        start = self.spStartRoll.value()
        end = self.spEndRoll.value()
        roll = [i for i in range(start, end)]
        records = []
        for i in range(len(data)):
            records.append([roll[i], data[i][0], float(data[i][1]), float(data[i][2]), float(data[i][3]), float(result[i])])
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Roll No.", "Name", "Exam 1", "Exam 2", "Practicals", "Result"])
            for record in records:
                ws.append(record)
            wb.save(file)
            self.status.showMessage("All records successfully saved to " + file, 2000)
        except Exception as e:
            QMessageBox.warning(self, "Error", str(e), QMessageBox.Ok)

    def setTableStudentUI(self):
        self.tableStudentHeader = self.tableStudent.horizontalHeader()

        self.tableStudentHeader.setSectionResizeMode(0, QHeaderView.Fixed)
        self.tableStudentHeader.setSectionResizeMode(1, QHeaderView.Stretch)
        self.tableStudentHeader.setSectionResizeMode(2, QHeaderView.Stretch)
        self.tableStudentHeader.setSectionResizeMode(3, QHeaderView.Stretch)
        self.tableStudent.setColumnWidth(0, 250)

        # self.tableStudent.setVisible(True)

    def setTableResultUI(self):
        self.tableResultHeader = self.tableResult.horizontalHeader()
        self.tableResultHeader.setSectionResizeMode(0, QHeaderView.Stretch)

    def compute(self):
        try:
            data = self.tableStudentModel.getAllRecords()
            exam1 = [float(row[1]) for row in data]
            exam2 = [float(row[2]) for row in data]
            practicals = [float(row[3]) for row in data]
            # print(exam1)
            # print(exam2)
            # print(practicals)
            result = Fuzzy.compute_fuzzy_performance(exam1, exam2, practicals)
            result = [str(x) for x in result]
            # print(result)
            self.tableResultModel.setAllRecords(result)
            self.tableResult.setModel(None)
            self.tableResult.setModel(self.tableResultModel)
            self.setTableResultUI()
        except ValueError:
            QMessageBox.warning(self, "Error", "Please make sure entire row is filled.", QMessageBox.Ok)
        except Exception as e:
            QMessageBox.warning(self, "Error", str(e), QMessageBox.Ok)
            # print(sys.exc_info()[0])

    def clear(self):
        self.tableStudentModel.clear()
        self.tableStudent.setModel(None)
        self.tableStudent.setModel(self.tableStudentModel)
        self.setTableStudentUI()

        self.tableResultModel.clear()
        self.tableResult.setModel(None)
        self.tableResult.setModel(self.tableResultModel)
        self.setTableResultUI()

    def updateRollNo(self):
        self.spStartRoll.setMaximum(self.spEndRoll.value())

        start = self.spStartRoll.value()
        end = self.spEndRoll.value()
        if (end-start) > 500:
            return

        self.tableStudentModel.updateRoll(start, end)
        self.tableResultModel.updateRoll(start, end)

        self.tableStudent.setModel(None)
        self.tableStudent.setModel(self.tableStudentModel)
        self.setTableStudentUI()

        self.tableResult.setModel(None)
        self.tableResult.setModel(self.tableResultModel)
        self.setTableResultUI()

    def setup_database(self):
        self.con = QSqlDatabase.addDatabase("QSQLITE")
        self.con.setDatabaseName("student_records.sqlite")

        if not self.con.open():
            QMessageBox.critical("Database Error!", self.con.lastError().databaseText())

    def guidelines(self):
        dlg = Guidelines.GuidelinesDialog()
        dlg.exec()

    def toggleTheme(self):
        if self.actionDark.isChecked():
            self.default_palette = QApplication.palette()
            dark_palette = QPalette()
            dark_palette.setColor(QPalette.Window, QColor(53, 53, 53))
            dark_palette.setColor(QPalette.WindowText, Qt.white)
            dark_palette.setColor(QPalette.Base, QColor(35, 35, 35))
            dark_palette.setColor(QPalette.AlternateBase, QColor(53, 53, 53))
            dark_palette.setColor(QPalette.ToolTipBase, QColor(25, 25, 25))
            dark_palette.setColor(QPalette.ToolTipText, Qt.white)
            dark_palette.setColor(QPalette.Text, Qt.white)
            dark_palette.setColor(QPalette.Button, QColor(53, 53, 53))
            dark_palette.setColor(QPalette.ButtonText, Qt.white)
            dark_palette.setColor(QPalette.BrightText, Qt.red)
            dark_palette.setColor(QPalette.Link, QColor(42, 130, 218))
            dark_palette.setColor(QPalette.Highlight, Qt.darkGray)
            dark_palette.setColor(QPalette.HighlightedText, QColor(35, 35, 35))
            dark_palette.setColor(QPalette.Active, QPalette.Button, QColor(53, 53, 53))
            dark_palette.setColor(QPalette.Disabled, QPalette.ButtonText, Qt.darkGray)
            dark_palette.setColor(QPalette.Disabled, QPalette.WindowText, Qt.darkGray)
            dark_palette.setColor(QPalette.Disabled, QPalette.Text, Qt.darkGray)
            dark_palette.setColor(QPalette.Disabled, QPalette.Light, QColor(53, 53, 53))
            app = QApplication.instance()
            if app is None:
                print("None")
                return
            app.setPalette(dark_palette)
        else:
            app = QApplication.instance()
            if app is None:
                print("None")
                return
            app.setPalette(self.default_palette)
            app.setStyleSheet(open("app_style.css").read())


def main():
    app = QApplication([])
    app.setStyleSheet(open("app_style.css").read())
    app.setStyle("Fusion")
    window = MainWindow()
    window.show()
    app.exec_()


if __name__ == "__main__":
    main()
